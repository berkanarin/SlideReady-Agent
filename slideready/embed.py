"""Slayt videolarını ORİJİNAL sunumun kendi slaytlarına gömer.

- Video slaytın en üst katmanında, tam ekran; orijinal şekiller altında düzenlenebilir kalır.
- Slayt açılınca otomatik başlar, döngü yoktur (isteyen PowerPoint'te Oynatma > Durdurulana kadar döngü'yü açar).
- Poster = videonun son karesi: oynamadan önce, PDF'te ve baskıda tamamlanmış slayt görünür.
- Slaytın eski giriş animasyonları kaldırılır (video artık onları içerir); slayt geçişleri korunur.
- Daha önce SlideReady'nin eklediği videolar yeniden çalıştırmada değiştirilir.
"""
import subprocess
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
PREFIX = "SlideReady Video"
# PowerPoint'in "Başlat: Otomatik" için yazdığı zamanlama (repeatCount yok = döngüsüz)
TIMING = f'''<p:timing xmlns:p="{NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{{DUR}}" fill="hold"/><p:tgtEl><p:spTgt spid="{{SPID}}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{{SPID}}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''


def duration(mp4):
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(mp4)],
                         capture_output=True, text=True, check=True).stdout.strip()
    return float(out or 0)


def last_frame(mp4, png):
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-sseof", "-0.1", "-i", str(mp4), "-update", "1",
                    "-frames:v", "1", str(png)], check=True)
    return png


def embed(deck, clips, dst):
    """clips: {slayt_no: mp4_yolu}. Döner: [(slayt_no, süre)]"""
    prs = Presentation(str(deck))
    done = []
    with tempfile.TemporaryDirectory() as tmp:
        for no, sl in enumerate(prs.slides, 1):
            mp4 = clips.get(no)
            if not mp4:
                continue
            for sh in list(sl.shapes):
                if sh.name.startswith(PREFIX):
                    sh._element.getparent().remove(sh._element)
            poster = last_frame(mp4, Path(tmp) / f"{no:03d}.png")
            mov = sl.shapes.add_movie(str(mp4), 0, 0, prs.slide_width, prs.slide_height,
                                      poster_frame_image=str(poster), mime_type="video/mp4")
            mov.name = f"{PREFIX} {no:02d}"
            old = sl._element.find(f"{{{NS}}}timing")
            if old is not None:
                sl._element.remove(old)
            dur = duration(mp4)
            timing = etree.fromstring(TIMING.replace("{DUR}", str(int(dur * 1000))).replace("{SPID}", str(mov.shape_id)))
            ext = sl._element.find(f"{{{NS}}}extLst")      # şema sırası: ... transition, timing, extLst
            if ext is not None:
                ext.addprevious(timing)
            else:
                sl._element.append(timing)
            done.append((no, round(dur, 2)))
        prs.save(str(dst))
    return done
