"""PPTX envanteri: her slayttaki şekillerin konumu, türü, metni ve verisi.

Koordinatlar sahne SVG birimindedir (1/100 mm = EMU / 360). Her şeklin kimliği (sid) PowerPoint'teki Shape.Id ile
aynıdır; katmanlar bu kimlikle çizilir ve sahnedeki grup kimliği "p<sid>" olur.
"""
import re
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from .media import kind as media_kind
from .regions import blocks as image_blocks

# Sayı: isteğe bağlı ön ek (₺, %, +), binlik ayraçlı ya da ondalıklı sayı, isteğe bağlı kısa son ek (%, M, bin, gün)
NUM = re.compile(r"^\s*([%₺$€£+\-−]{0,2})\s*(\d{1,3}(?:[.,]\d{3})+(?:[.,]\d+)?|\d+(?:[.,]\d+)?)\s*([%+]|[^\W\d_]{1,6}\.?)?\s*$")
PIE = {XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.DOUGHNUT_EXPLODED, XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED,
       XL_CHART_TYPE.THREE_D_PIE, XL_CHART_TYPE.THREE_D_PIE_EXPLODED}
U = 360  # EMU -> 1/100 mm


def number_parts(text):
    """'₺12,4M' -> ('₺', '12,4', 'M'); sayı değilse ya da sıra etiketi/yılsa None."""
    m = NUM.match(text)
    if not m:
        return None
    pre, val, suf = m[1], m[2], m[3] or ""
    if re.fullmatch(r"0\d+", val):                      # 01, 02: sıra etiketi
        return None
    if not pre and not suf and re.fullmatch(r"(19|20)\d\d", val):  # 2026: yıl
        return None
    return pre, val, suf


def chart_kind(chart):
    ct = chart.chart_type
    if ct in PIE:
        return "pie"
    name = str(ct)
    return "bar_h" if "BAR" in name else "line" if ("LINE" in name or "AREA" in name or "XY" in name) else "column"


def _shape(sh, W, H, off=(0, 0, 1, 1)):
    ox, oy, sx, sy = off
    d = dict(name=sh.name, sid=sh.shape_id, id=f"p{sh.shape_id}", x=ox + sh.left * sx / U, y=oy + sh.top * sy / U, w=sh.width * sx / U, h=sh.height * sy / U)
    text = sh.text_frame.text.strip() if sh.has_text_frame else ""
    if getattr(sh, "has_chart", False) and sh.has_chart:
        c = sh.chart
        series = [dict(name=s.name, values=[float(v or 0) for v in s.values]) for p in c.plots for s in p.series]
        try:
            cats = [str(x) for x in c.plots[0].categories]
        except Exception:
            cats = []
        d.update(kind="chart", chart=chart_kind(c), series=series, categories=cats)
    elif getattr(sh, "has_table", False) and sh.has_table:
        t = sh.table
        d.update(kind="table", rows=len(t.rows), cols=len(t.columns),
                 cells=[[t.cell(r, c).text for c in range(len(t.columns))] for r in range(len(t.rows))])
    elif media_kind(sh):
        d.update(kind="media")                          # GIF/video: kendi kareleriyle oynar, giriş hareketi yok
    elif text and number_parts(text):
        pre, val, suf = number_parts(text)
        d.update(kind="number", text=text, pre=pre, val=val, suf=suf)
    elif text:
        d.update(kind="text", text=text[:200], words=len(text.split()))
    elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        d.update(kind="picture")
        crop = any(getattr(sh, "crop_" + k, 0) for k in ("left", "top", "right", "bottom"))
        if sh.width * sh.height > 0.08 * W * H and not crop:
            try:
                reg = image_blocks(sh.image.blob)   # resim içindeki metin/şekil blokları (ekran görüntüsü vb.)
            except Exception:
                reg = None
            if reg:
                d["regions"] = reg
    else:
        area = sh.width * sh.height / (W * H)
        d.update(kind="panel" if area > 0.03 else "line" if min(sh.width, sh.height) < 20000 else "deco")
    if sh.is_placeholder:
        pt = sh.placeholder_format.type
        if pt in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE):
            d["kind"] = "title"
        elif pt in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE):
            d["kind"] = "static"
    return d


def _walk(shapes, W, H, off=(0, 0, 1, 1), group=None):
    for sh in shapes:
        if sh.width is None:
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            # grup içi koordinatlar chOff/chExt uzayında: mutlak konuma çevir
            xfrm = sh._element.grpSpPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
            ch_off, ch_ext = xfrm.find("{*}chOff"), xfrm.find("{*}chExt")
            cx, cy = int(ch_off.get("x")), int(ch_off.get("y"))
            cw, chh = max(int(ch_ext.get("cx")), 1), max(int(ch_ext.get("cy")), 1)
            ox, oy, sx, sy = off
            gsx, gsy = sh.width / cw, sh.height / chh
            gx, gy = ox + sh.left * sx / U, oy + sh.top * sy / U
            sub = (gx - cx * gsx * sx / U, gy - cy * gsy * sy / U, sx * gsx, sy * gsy)
            yield from _walk(sh.shapes, W, H, sub, group or sh.name)
            continue
        d = _shape(sh, W, H, off)
        if group:
            d["group"] = group
        yield d


def _inside(o, d):
    cx, cy = o["x"] + o["w"] / 2, o["y"] + o["h"] / 2
    return d["x"] <= cx <= d["x"] + d["w"] and d["y"] <= cy <= d["y"] + d["h"]


def fonts(prs):
    """Sunumda kullanılan yazı tipi aileleri (tema fontları çözülerek)."""
    names = set()
    theme = {}
    try:
        master = prs.slide_masters[0]
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                xml = rel.target_part.blob.decode("utf-8", "ignore")
                maj = re.search(r"<a:majorFont>\s*<a:latin typeface=\"([^\"]*)\"", xml)
                mnr = re.search(r"<a:minorFont>\s*<a:latin typeface=\"([^\"]*)\"", xml)
                theme = {"+mj-lt": maj[1] if maj else "", "+mn-lt": mnr[1] if mnr else ""}
    except Exception:
        pass
    parts = [s.part for s in prs.slides] + [l.part for l in prs.slide_layouts] + [m.part for m in prs.slide_masters]
    for p in parts:
        for f in re.findall(r'<a:latin typeface="([^"]+)"', p.blob.decode("utf-8", "ignore")):
            f = theme.get(f, f)
            if f and not f.startswith("+"):
                names.add(f)
    names.update(v for v in theme.values() if v)
    return sorted(names)


def inventory(path):
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    slides = []
    for i, sl in enumerate(prs.slides, 1):
        shapes = list(_walk(sl.shapes, W, H))
        for d in shapes:                # üstünde başka şekil olan resim arka plandır: bloklara bölünmez
            if "regions" in d and any(o is not d and _inside(o, d) for o in shapes):
                del d["regions"]
        if shapes and not any(d["kind"] == "title" for d in shapes):
            top = min((d for d in shapes if d["kind"] == "text"), key=lambda d: (d["y"], d["x"]), default=None)
            if top and top["y"] < H / U * 0.2:
                top["kind"] = "title"
        hidden = sl._element.get("show") == "0"
        slides.append(dict(no=i, hidden=hidden, shapes=shapes))
    return dict(width=W / U, height=H / U, slides=slides, fonts=fonts(prs))
