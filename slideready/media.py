"""Hareketli medya: animasyonlu GIF'ler ve videolar (slaytta, düzende ya da asıl slaytta).

PowerPoint katmanlarında bunlar tek bir durağan kare olarak görünür. Bu modül medyayı bulur, kareleri ffmpeg ile
çıkarır; sahne her anda doğru kareyi gösterir. Böylece
arka plan videosu ya da GIF, verilerle birlikte tek bir bütün olarak oynar.
"""
import io
import json
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

FPS = 30            # medya kare hızı (sahne bu kareleri zamana göre seçer)
MAX_SECONDS = 60    # bir video slayt süresini en fazla bu kadar uzatır
U = 360             # EMU -> 1/100 mm
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
P14 = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"


def _video_part(sh):
    """Video şeklinin gömülü medya parçası (p14:media, yoksa a:videoFile); dış bağlantıysa None."""
    nvpr = sh._element.find(f"{P}nvPicPr/{P}nvPr")
    if nvpr is None:
        return None
    rids = [e.get(f"{R}embed") for e in nvpr.iter(f"{P14}media")]
    rids += [e.get(f"{R}link") for e in nvpr.iter(f"{A}videoFile")]
    for rid in filter(None, rids):
        rel = sh.part.rels.get(rid)
        if rel is not None and not rel.is_external:
            return rel.target_part
    return None


def _is_video(sh):
    return sh.shape_type == MSO_SHAPE_TYPE.MEDIA or sh._element.find(f".//{A}videoFile") is not None


def _gif_frames(blob):
    try:
        return getattr(Image.open(io.BytesIO(blob)), "n_frames", 1)
    except Exception:
        return 1


def kind(sh):
    """'video', 'gif' (birden çok kareli) ya da None."""
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        return None
    if _is_video(sh):
        return "video" if _video_part(sh) is not None else None
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            img = sh.image
        except Exception:
            return None
        if img.content_type == "image/gif" and _gif_frames(img.blob) > 1:
            return "gif"
    return None


def _walk(shapes, off=(0, 0, 1, 1)):
    """(şekil, mutlak kutu) çiftleri; gruplar açılır (deck._walk ile aynı dönüşüm)."""
    for sh in shapes:
        if sh.width is None:
            continue
        ox, oy, sx, sy = off
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = sh._element.grpSpPr.find(f"{A}xfrm")
            ch_off, ch_ext = xfrm.find(f"{A}chOff"), xfrm.find(f"{A}chExt")
            cx, cy = int(ch_off.get("x")), int(ch_off.get("y"))
            gsx = sh.width / max(int(ch_ext.get("cx")), 1)
            gsy = sh.height / max(int(ch_ext.get("cy")), 1)
            gx, gy = ox + sh.left * sx / U, oy + sh.top * sy / U
            yield from _walk(sh.shapes, (gx - cx * gsx * sx / U, gy - cy * gsy * sy / U, sx * gsx, sy * gsy))
            continue
        yield sh, (ox + sh.left * sx / U, oy + sh.top * sy / U, sh.width * sx / U, sh.height * sy / U)


def _shows_master(el):
    return el.get("showMasterSp") not in ("0", "false")


def _layers(slide):
    """Slaytta görünen katmanlar, alttan üste: asıl slayt, düzen, slayt (showMasterSp'ye uyarak)."""
    layout = slide.slide_layout
    master = layout.slide_master
    out = [(slide, "slide")]
    if _shows_master(slide._element):
        out.insert(0, (layout, "layout"))
        if _shows_master(layout._element):
            out.insert(0, (master, "master"))
    return out


def _probe(path):
    r = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration:stream=codec_type", "-of", "json",
                        str(path)], capture_output=True, text=True)
    try:
        info = json.loads(r.stdout or "{}")
    except ValueError:
        info = {}
    dur = float(info.get("format", {}).get("duration") or 0)
    audio = any(s.get("codec_type") == "audio" for s in info.get("streams", []))
    return dur, audio


def collect(deck, work):
    """Her slayt için hareketli medya listesi; medya dosyaları work/media/ altına bir kez yazılır.

    Döner: {slayt_no: [dict(type, name, sid, layer, x, y, w, h, file, duration, audio)]}
    """
    prs = Presentation(str(deck))
    mdir = work / "media"
    files, out = {}, {}
    for no, sl in enumerate(prs.slides, 1):
        items = []
        for owner, layer in _layers(sl):
            for sh, (x, y, w, h) in _walk(owner.shapes):
                if layer != "slide" and sh.is_placeholder:
                    continue           # düzen/asıl slayt yer tutucuları slaytta görünmez
                k = kind(sh)
                if not k:
                    continue
                part = _video_part(sh) if k == "video" else sh.image
                key = str(part.partname) if k == "video" else part.sha1
                if key not in files:
                    ext = Path(str(part.partname)).suffix if k == "video" else ".gif"
                    mdir.mkdir(parents=True, exist_ok=True)
                    f = mdir / f"m{len(files) + 1:02d}{ext}"
                    if not f.exists() or f.stat().st_size != len(part.blob):
                        f.write_bytes(part.blob)
                    dur, audio = _probe(f) if k == "video" else (0.0, False)
                    files[key] = dict(file=str(f), duration=round(dur, 2), audio=audio)
                items.append(dict(type=k, name=sh.name, sid=sh.shape_id, layer=layer, x=x, y=y, w=w, h=h, **files[key]))
        if items:
            out[no] = items
    return out


def frames(item, width=None):
    """Medyanın karelerini (FPS) çıkarır; kare klasörü ve kare sayısını döndürür. Sonuç önbelleğe alınır."""
    src = Path(item["file"])
    ext = "png" if item["type"] == "gif" else "jpg"
    scale = f",scale='min({int(width)},iw)':-2" if width and item["type"] == "video" else ""
    d = src.with_name(src.stem + (f"-{int(width)}" if scale else ""))
    stamp = d / "done.txt"
    if not stamp.exists():
        d.mkdir(parents=True, exist_ok=True)
        for old in d.glob("f*.*"):
            old.unlink()
        args = ["-t", str(MAX_SECONDS)] if item["type"] == "video" else []
        q = ["-q:v", "3"] if ext == "jpg" else []
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", *args, "-i", str(src), "-vf", f"fps={FPS}{scale}", *q,
                        str(d / f"f%05d.{ext}")], check=True)
        stamp.write_text(str(len(list(d.glob(f"f*.{ext}")))), encoding="utf-8")
    return d, int(stamp.read_text(encoding="utf-8")), ext
