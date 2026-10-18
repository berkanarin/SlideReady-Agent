"""Teslim öncesi kalite kontrolü: render edilmiş klipleri değişmez kurallara göre ölçer.

  slideready check SUNUM.pptx           önizleme ve çıktı kliplerini, varsa çıktı PPTX'in posterlerini denetler

Kurallar (SKILL.md "Değişmez kurallar"):
  1. İlk kare boş: açılıştaki en az 0,3 sn yalnızca arka plandır (içerik pikselleri arka plana yakın).
  2. Efektler tekrar etmez: plan.json'daki kahraman efekt ve açılış imzası slayttan slayta farklıdır; görsel açılış
     imzası (içeriğin hangi yönden, hangi sırayla geldiği ve kameranın ne kadar oynadığı) art arda iki slaytta aynı olmaz.
  3. Son kare PowerPoint'in kendi çizimidir (full.png).
  4. Okuma payında hiçbir şey kımıldamaz.
  5. Kare kaçmaz: kare sayısı süre x fps, sabit kare hızı, hareket ortasında yinelenen (takılan) kare yok.

Ölçümler küçültülmüş karelerde yapılır; hareketli medya (arka plan videosu, GIF) kutuları ölçümün dışında tutulur.
"""
import io
import json
import math
import subprocess
from pathlib import Path

import numpy as np
from PIL import Image

SIZE = (192, 108)        # ölçüm çözünürlüğü
OPEN_SECONDS = 0.3       # boş açılışın en kısa süresi
INK = 25                 # bir pikseli "içerik" sayan fark (full ile bg arası, 0-255)
EMPTY_FRAC = 0.02        # boş sayılan karede tam slayta yakın içerik pikseli oranı en çok
LAST_DIFF = 2.0          # son kare ile full.png arası ortalama fark en çok (0-255, sıkıştırma payıyla)
STILL = 0.12             # iki kare arası "kımıldamıyor" sayılan ortalama fark
MOVING = 0.6             # iki kare arası "hareket var" sayılan ortalama fark
SIMILAR = 0.12           # iki açılış imzası arası bu uzaklığın altındaysa aynı hareket sayılır


def _probe(mp4):
    out = subprocess.run(["ffprobe", "-v", "error", "-count_frames", "-select_streams", "v:0", "-show_entries",
                          "stream=nb_read_frames,r_frame_rate,avg_frame_rate:format=duration", "-of", "json", str(mp4)],
                         capture_output=True, text=True, check=True).stdout
    info = json.loads(out)
    st = info["streams"][0]

    def rate(s):
        a, _, b = s.partition("/")
        return float(a) / float(b or 1) if float(b or 1) else 0.0

    return dict(frames=int(st.get("nb_read_frames") or 0), fps=rate(st["r_frame_rate"]), avg=rate(st["avg_frame_rate"]),
                duration=float(info["format"].get("duration") or 0))


def frames(mp4, size=SIZE):
    """Bütün kareler (n, h, w, 3) float32, 0-255."""
    w, h = size
    raw = subprocess.run(["ffmpeg", "-v", "error", "-i", str(mp4), "-vf", f"scale={w}:{h}:flags=area", "-f", "rawvideo",
                          "-pix_fmt", "rgb24", "-"], capture_output=True, check=True).stdout
    return np.frombuffer(raw, np.uint8).reshape(-1, h, w, 3).astype(np.float32)


def image(path, size=SIZE):
    """PNG, kliplerle aynı yoldan (yuv420p'ye çevrilip küçültülerek) okunur; renk dönüşümü fark sayılmasın."""
    w, h = size
    raw = subprocess.run(["ffmpeg", "-v", "error", "-i", str(path), "-vf", f"format=yuv420p,scale={w}:{h}:flags=area",
                          "-f", "rawvideo", "-pix_fmt", "rgb24", "-frames:v", "1", "-"], capture_output=True).stdout
    if len(raw) == w * h * 3:
        return np.frombuffer(raw, np.uint8).reshape(h, w, 3).astype(np.float32)
    return np.asarray(Image.open(path).convert("RGB").resize(size, Image.BOX), dtype=np.float32)


def media_mask(media, slide_w, slide_h, size=SIZE):
    """Hareketli medya kutuları dışında kalan pikseller True."""
    w, h = size
    keep = np.ones((h, w), bool)
    for m in media or []:
        x0, y0 = int(m["x"] / slide_w * w), int(m["y"] / slide_h * h)
        x1, y1 = math.ceil((m["x"] + m["w"]) / slide_w * w), math.ceil((m["y"] + m["h"]) / slide_h * h)
        keep[max(0, y0):max(0, y1), max(0, x0):max(0, x1)] = False
    return keep


def content_mask(full, bg, keep):
    return (np.abs(full - bg).max(axis=2) > INK) & keep


def filled(frame, full, bg, mask):
    """Karede tam slayta arka plandan daha yakın duran içerik piksellerinin oranı (0 = boş, 1 = dolu)."""
    if not mask.any():
        return 0.0
    near_full = np.abs(frame - full).sum(axis=2) < np.abs(frame - bg).sum(axis=2)
    return float((near_full & mask).sum() / mask.sum())


def steps(fr, keep):
    """Ardışık kareler arası ortalama fark (medya dışında); steps[i] = fark(i-1, i), steps[0] = 0."""
    d = np.abs(np.diff(fr, axis=0)).mean(axis=3)
    d = d[:, keep].mean(axis=1)
    return np.concatenate([[0.0], d])


def stutters(st):
    """Hareketin ortasında yinelenen kareler: önce ve sonra hareket var, arada kare aynı kalmış."""
    return [i for i in range(1, len(st) - 1) if st[i] < STILL and st[i - 1] > MOVING and st[i + 1] > MOVING]


def signature(fr, full, bg, mask, keep, fps):
    """Açılışın nasıl kurulduğu (içerikten bağımsız): içerik hangi yönden ve sırayla geliyor, ne kadar sürede, kamera
    ne kadar oynuyor. Döner: [x yönü, y yönü, merkezden dışa, açılış süresi, kamera] ya da içerik yoksa None."""
    if mask.sum() < 50:
        return None
    n = len(fr)
    near = np.abs(fr - full[None]).sum(axis=3) < np.abs(fr - bg[None]).sum(axis=3)       # (n, h, w)
    arrived = np.where(near.any(axis=0), near.argmax(axis=0), n)                       # ilk "geldiği" kare
    ys, xs = np.nonzero(mask)
    t = arrived[ys, xs].astype(np.float64)
    if t.std() < 1e-6:
        return None
    h, w = mask.shape
    r = np.hypot(xs / w - 0.5, (ys / h - 0.5) * h / w)

    def corr(v):
        v = v.astype(np.float64)
        return float(np.corrcoef(v, t)[0, 1]) if v.std() > 1e-9 else 0.0

    span = float(np.percentile(t, 90) - np.percentile(t, 10)) / fps
    bgpix = keep & ~mask
    cam = 0.0
    if bgpix.any():
        early = fr[: min(n, int(fps * 1.5))]
        cam = float(np.abs(early - bg[None])[:, bgpix].mean()) / 255
    return [round(corr(xs), 3), round(corr(ys), 3), round(corr(r), 3), round(min(span / 3, 1), 3), round(min(cam * 8, 1), 3)]


def distance(a, b):
    return float(np.linalg.norm(np.subtract(a, b)) / math.sqrt(len(a)))


def check_clip(mp4, layer_dir, media, slide_w, slide_h, still_at=None):
    """Tek klibin ölçümü. Döner: dict(errors=[...], warnings=[...], sig=..., ...)"""
    errors, warnings = [], []
    p = _probe(mp4)
    fps = p["fps"] or 30
    expected = round(p["duration"] * fps)
    if abs(p["frames"] - expected) > 1:
        errors.append(f"kare sayısı {p['frames']}, beklenen {expected} ({p['duration']:.2f} sn x {fps:g} fps): kare kaçıyor")
    if p["avg"] and abs(p["avg"] - fps) > 0.01 * fps:
        errors.append(f"kare hızı sabit değil (ortalama {p['avg']:.2f}, beklenen {fps:g})")
    fr = frames(mp4)
    full, bg = image(layer_dir / "full.png"), image(layer_dir / "bg.png")
    keep = media_mask(media, slide_w, slide_h)
    mask = content_mask(full, bg, keep)
    out = dict(frames=p["frames"], fps=fps, duration=round(p["duration"], 2))
    if len(fr):
        first = filled(fr[0], full, bg, mask)
        out["first_filled"] = round(first, 3)
        if first > EMPTY_FRAC:
            errors.append(f"ilk kare boş değil: içeriğin %{first * 100:.0f}'i görünüyor")
        k = max(1, math.ceil(OPEN_SECONDS * fps))
        late = next((i for i in range(min(k, len(fr))) if filled(fr[i], full, bg, mask) > EMPTY_FRAC), None)
        if late is not None and first <= EMPTY_FRAC:
            errors.append(f"boş açılış {late / fps:.2f} sn sürüyor, en az {OPEN_SECONDS} sn olmalı "
                          "(PowerPoint videoyu başlatırken ilk kareleri yutabilir)")
        last = float(np.abs(fr[-1] - full)[keep].mean()) if keep.any() else 0.0
        out["last_diff"] = round(last, 2)
        if last > LAST_DIFF:
            errors.append(f"son kare PowerPoint çiziminden farklı (ortalama fark {last:.1f}/255)")
        if keep.mean() < 0.1:
            warnings.append("slaytın neredeyse tamamı hareketli medya: takılma ve okuma payı ölçülemedi, gözle kontrol edin")
            out["sig"] = signature(fr, full, bg, mask, keep, fps)
            out.update(errors=errors, warnings=warnings)
            return out
        if not mask.any():
            warnings.append("arka plandan ayrılan içerik yok: ilk kare ve açılış imzası ölçülemedi")
        st = steps(fr, keep)
        bad = stutters(st)
        if bad:
            errors.append(f"hareket ortasında {len(bad)} takılan kare (ilk: {bad[0] / fps:.2f} sn)")
        if still_at is not None:
            i0 = min(len(st), int(math.ceil(still_at * fps)) + 1)
            moving = [i for i in range(i0, len(st)) if st[i] > STILL]
            if moving:
                errors.append(f"okuma payında hareket var ({moving[0] / fps:.2f} sn'de)")
        out["sig"] = signature(fr, full, bg, mask, keep, fps)
    out.update(errors=errors, warnings=warnings)
    return out


REUSE_GAP = 4            # bir kahraman koreografi ancak arada en az bu kadar canlı slayt varsa yeniden kullanılabilir


def check_plan(plans):
    """plan.json düzeyinde tekrar. Döner: (hatalar, uyarılar)"""
    errors, warnings = [], []
    live = [p for p in plans if p.get("animate", True)]
    seen = {}
    for i, p in enumerate(live):
        hero = p.get("hero")
        if not hero:
            continue
        if hero in seen:
            j, no = seen[hero]
            msg = f"slayt {p['no']}: kahraman koreografi '{hero}' slayt {no}'de de kullanılmış"
            (warnings if i - j > REUSE_GAP else errors).append(msg)
        seen[hero] = (i, p["no"])

    def key(p):
        return (p.get("hero"), p.get("chosen") or p.get("style"), round(float(p.get("zoom") or 1), 2), p.get("lead"))

    for a, b in zip(live, live[1:]):
        if key(a) == key(b):
            errors.append(f"slayt {a['no']} ve {b['no']}: aynı açılış (koreografi, tarz, kamera, öne çıkan aynı)")
    if len(live) >= 3 and not any(p.get("hero") for p in live):
        warnings.append("hiçbir slayta kahraman koreografi (hero) yazılmamış: Yönetmen adımı atlanmış")
    return errors, warnings


def check_signatures(rows):
    """Görsel açılış imzaları: ardışık iki slayt aynı biçimde açılıyorsa hata, deste içinde tekrar ediyorsa uyarı."""
    errors, warnings = [], []
    sigs = [(no, r["sig"]) for no, r in rows if r.get("sig")]
    for (a, sa), (b, sb) in zip(sigs, sigs[1:]):
        if b - a == 1 and distance(sa, sb) < SIMILAR:
            errors.append(f"slayt {a} ve {b} aynı biçimde açılıyor (görsel imza uzaklığı {distance(sa, sb):.2f})")
    for i, (a, sa) in enumerate(sigs):
        for b, sb in sigs[i + 2:]:
            if distance(sa, sb) < SIMILAR / 2:
                warnings.append(f"slayt {a} ve {b} neredeyse aynı biçimde açılıyor")
    return errors, warnings


def posters(pptx, nos, work, inv):
    """Çıktı PPTX'teki SlideReady videolarının posterleri boş açılış karesi mi, videolar döngüsüz mü?"""
    from pptx import Presentation
    from lxml import etree

    errors = []
    prs = Presentation(str(pptx))
    for no, sl in enumerate(prs.slides, 1):
        if no not in nos:
            continue
        vids = [sh for sh in sl.shapes if sh.name.startswith("SlideReady Video")]
        if not vids:
            errors.append(f"slayt {no}: çıktıda SlideReady videosu yok")
            continue
        if b'repeatCount="indefinite"' in etree.tostring(sl._element):
            errors.append(f"slayt {no}: video döngüde")
        rid = vids[0]._element.xpath(".//a:blip/@r:embed")
        if not rid:
            errors.append(f"slayt {no}: videonun posteri yok")
            continue
        img = np.asarray(Image.open(io.BytesIO(sl.part.related_part(rid[0]).blob)).convert("RGB").resize(SIZE, Image.BOX),
                         dtype=np.float32)
        d = work / "layers" / f"{no:03d}"
        full, bg = image(d / "full.png"), image(d / "bg.png")
        keep = media_mask(inv["slides"][no - 1].get("media"), inv["width"], inv["height"])
        f = filled(img, full, bg, content_mask(full, bg, keep))
        if f > EMPTY_FRAC:
            errors.append(f"slayt {no}: poster boş değil (içeriğin %{f * 100:.0f}'i görünüyor)")
    return errors


def run(deck, work, pptx=None):
    """Bütün denetimler. Döner: (rapor dict, başarılı mı)"""
    inv = json.loads((work / "analysis.json").read_text(encoding="utf-8"))
    plans = json.loads((work / "plan.json").read_text(encoding="utf-8"))["slides"]
    perr, pwarn = check_plan(plans)
    report = dict(plan=perr, plan_warnings=pwarn, clips={}, sets={})
    ok = not report["plan"]
    for kind in ("clips", "preview"):
        rows = []
        for mp4 in sorted((work / kind).glob("[0-9][0-9][0-9].mp4")):
            no = int(mp4.stem)
            d = work / "layers" / f"{no:03d}"
            if not (d / "full.png").exists():
                continue
            scene = work / "scenes" / f"{no:03d}" / "scene.json"
            still = json.loads(scene.read_text(encoding="utf-8"))["style"].get("still") if scene.exists() else None
            s = inv["slides"][no - 1]
            rows.append((no, check_clip(mp4, d, s.get("media"), inv["width"], inv["height"], still)))
        if not rows:
            continue
        errs, warns = check_signatures(rows)
        report["sets"][kind] = dict(errors=errs, warnings=warns)
        report["clips"][kind] = {no: r for no, r in rows}
        ok = ok and not errs and all(not r["errors"] for _, r in rows)
    if pptx and Path(pptx).exists() and report["clips"].get("clips"):
        report["posters"] = posters(pptx, set(report["clips"]["clips"]), work, inv)
        ok = ok and not report["posters"]
    (work / "qa.json").write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    return report, ok


def show(report, ok):
    lines = []
    lines += [f"✗ plan: {e}" for e in report["plan"]] + [f"! plan: {w}" for w in report.get("plan_warnings", [])]
    for kind, rows in report["clips"].items():
        lines.append(f"{'çıktı' if kind == 'clips' else 'önizleme'}:")
        for no, r in rows.items():
            mark = "✗" if r["errors"] else "✓"
            info = f"{r['duration']:.1f} sn, {r['frames']} kare, ilk kare dolu %{r.get('first_filled', 0) * 100:.0f}, " \
                   f"son kare farkı {r.get('last_diff', 0):.1f}/255"
            lines.append(f"  {mark} slayt {no}: {info}")
            lines += [f"      ✗ {e}" for e in r["errors"]] + [f"      ! {w}" for w in r["warnings"]]
        s = report["sets"].get(kind, {})
        lines += [f"  ✗ {e}" for e in s.get("errors", [])] + [f"  ! {w}" for w in s.get("warnings", [])]
    for e in report.get("posters", []):
        lines.append(f"✗ poster: {e}")
    if not report["clips"]:
        lines.append("! denetlenecek klip yok: önce preview ya da export çalıştırın")
    lines.append("✓ Değişmez kurallar sağlanıyor." if ok and report["clips"] else
                 "✗ TESLİM ETME: yukarıdaki sorunları düzeltin.")
    print("\n".join(lines))
