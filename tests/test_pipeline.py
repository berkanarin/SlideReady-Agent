"""Hızlı testler: python -m unittest discover tests

Çizim testleri PowerPoint (Windows, COM) gerektirir ve başka yerde atlanır; geri kalanı saf Python'dır."""
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from slideready.deck import inventory, number_parts  # noqa: E402
from slideready.media import collect  # noqa: E402
from slideready.pptrender import FRAMES, available, count_frames, matte, number_format, render  # noqa: E402
from slideready.regions import blocks  # noqa: E402

DEMO = ROOT / "examples" / "demo.pptx"
POWERPOINT = available()[0]


def run(*args):
    return subprocess.run([sys.executable, "-m", "slideready", *map(str, args)], cwd=ROOT, check=True,
                          capture_output=True)


class Numbers(unittest.TestCase):
    def test_counters(self):
        self.assertEqual(number_parts("₺12,4M"), ("₺", "12,4", "M"))
        self.assertEqual(number_parts("%38"), ("%", "38", ""))
        self.assertEqual(number_parts("1.250"), ("", "1.250", ""))
        self.assertEqual(number_parts("+%31"), ("+%", "31", ""))
        self.assertEqual(number_parts("9 gün"), ("", "9", "gün"))

    def test_not_counters(self):
        for s in ("01", "2026", "2026 Yıl Sonu Özeti", "Başlık"):
            self.assertIsNone(number_parts(s), s)


@unittest.skipUnless(POWERPOINT, "PowerPoint yok (yalnızca Windows)")
class Layers(unittest.TestCase):
    def test_every_shape_has_a_layer_and_matches_powerpoint(self):
        with tempfile.TemporaryDirectory() as tmp:
            run("analyze", DEMO, "-w", tmp)
            inv = inventory(DEMO)
            nos = [s["no"] for s in inv["slides"]]
            deck = render(DEMO, inv, Path(tmp), nos, width=960, log=lambda *_: None)
            for s in inv["slides"]:
                info = deck.info[s["no"]]
                self.assertEqual(info["missing"], [], f"slayt {s['no']}")
                self.assertEqual(sorted(info["layers"]), sorted(str(d["sid"]) for d in s["shapes"]), f"slayt {s['no']}")
                self.assertLess(info["fidelity"]["mean"], 1, f"slayt {s['no']}")
                body = deck.slide(s["no"])["body"]
                self.assertTrue(all(f'id="{d["id"]}"' in body for d in s["shapes"]))
            # sayaç kareleri: 4 sayı şekli, her biri PowerPoint'in çizdiği ara değerlerle
            nums = [d for d in inv["slides"][0]["shapes"] if d["kind"] == "number"]
            self.assertEqual(len(nums), 4)
            for d in nums:
                self.assertEqual(len(deck.info[1]["layers"][str(d["sid"])]["frames"]), FRAMES - 1)


class Analyze(unittest.TestCase):
    def test_analyze_writes_plan(self):
        """analyze çizim yapmaz (PowerPoint gerekmez); şekil kimlikleri slayt içinde tekildir."""
        with tempfile.TemporaryDirectory() as tmp:
            run("analyze", DEMO, "-w", tmp)
            self.assertFalse((Path(tmp) / "layers").exists())
            plan = json.loads((Path(tmp) / "plan.json").read_text(encoding="utf-8"))
            self.assertEqual(len(plan["slides"]), 4)
            self.assertTrue(all(p["animate"] for p in plan["slides"]))
            inv = json.loads((Path(tmp) / "analysis.json").read_text(encoding="utf-8"))
            ids = [d["id"] for d in inv["slides"][0]["shapes"]]
            self.assertEqual(len(ids), len(set(ids)))


class Matte(unittest.TestCase):
    """Fark matlama: beyaz ve siyah zemin üzerindeki iki kareden renk ve saydamlık geri kazanılır."""

    def test_recovers_color_and_alpha(self):
        from PIL import Image
        import numpy as np
        rnd = np.random.default_rng(1)
        color = rnd.integers(0, 256, (40, 60, 3)).astype(np.float32)
        alpha = np.zeros((40, 60), np.float32)
        alpha[10:30, 5:50] = rnd.random((20, 45))                  # yarı saydam gölge gibi
        alpha[15:25, 20:30] = 1

        def over(bg):
            a = alpha[..., None]
            return Image.fromarray((color * a + bg * (1 - a) + 0.5).astype(np.uint8))
        im, box = matte(over(255.0), over(0.0))
        self.assertEqual(box, (5, 10, 50, 30))
        rgba = np.asarray(im, dtype=np.float32)
        self.assertLess(np.abs(rgba[..., 3] / 255 - alpha).max(), 2 / 255)
        a = alpha[..., None]
        back = rgba[..., :3] * rgba[..., 3:] / 255 + 128 * (1 - rgba[..., 3:] / 255)
        self.assertLess(np.abs(back - (color * a + 128 * (1 - a))).max(), 2)

    def test_empty_layer(self):
        from PIL import Image
        _, box = matte(Image.new("RGB", (8, 8), "white"), Image.new("RGB", (8, 8), "black"))
        self.assertIsNone(box)


class CounterFrames(unittest.TestCase):
    def test_keeps_number_format(self):
        self.assertEqual(count_frames("12,4", 5), ["0,0", "3,1", "6,2", "9,3"])
        self.assertEqual(count_frames("1.250", 3), ["0", "625"])
        self.assertEqual(count_frames("3.250,75", 2), ["0,00"])
        self.assertEqual(number_format("1.250")(1), "1.250")
        self.assertEqual(number_format("12.500.000")(0.5), "6.250.000")


class ImageRegions(unittest.TestCase):
    """Metin/şekil resim olarak konmuşsa resim okuma sırasıyla bloklara ayrılır; fotoğraf bölünmez."""

    @staticmethod
    def png(im):
        import io
        buf = io.BytesIO()
        im.save(buf, "PNG")
        return buf.getvalue()

    def test_screenshot_splits_in_reading_order(self):
        from PIL import Image, ImageDraw
        im = Image.new("RGB", (1600, 900), (245, 247, 250))
        d = ImageDraw.Draw(im)
        d.rectangle((80, 60, 700, 140), fill=(20, 30, 90))                 # başlık
        for i in range(3):                                                  # üç kart
            d.rectangle((80 + i * 500, 260, 500 + i * 500, 520), fill=(20, 60, 180))
        d.rectangle((80, 640, 1500, 820), fill=(240, 170, 0))               # alt şerit
        r = blocks(self.png(im))
        self.assertEqual(r["bg"], "#f5f7fa")
        self.assertEqual(len(r["blocks"]), 5)
        xs = [b[0] for b in r["blocks"][1:4]]
        self.assertEqual(xs, sorted(xs))
        self.assertLess(r["blocks"][0][1], r["blocks"][1][1])

    def test_transparent_background(self):
        from PIL import Image, ImageDraw
        im = Image.new("RGBA", (800, 400), (0, 0, 0, 0))
        d = ImageDraw.Draw(im)
        d.rectangle((40, 40, 300, 360), fill=(200, 0, 0, 255))
        d.rectangle((460, 40, 760, 360), fill=(0, 0, 200, 255))
        r = blocks(self.png(im))
        self.assertIsNone(r["bg"])
        self.assertEqual(len(r["blocks"]), 2)

    def test_photo_is_not_split(self):
        import random
        from PIL import Image
        rnd = random.Random(1)
        im = Image.new("RGB", (320, 180))
        im.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(320 * 180)])
        self.assertIsNone(blocks(self.png(im)))


def media_deck(tmp):
    """Kurgusal deste: 1. slaytta tam ekran arka plan videosu + üstünde veri, düzende animasyonlu bir GIF."""
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Emu
    tmp = Path(tmp)
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc2=size=640x360:rate=30:duration=3",
                    "-f", "lavfi", "-i", "sine=frequency=440:duration=3", "-pix_fmt", "yuv420p", "-shortest",
                    str(tmp / "bg.mp4")], check=True)
    frames = []
    for i in range(10):
        im = Image.new("RGBA", (100, 100), (0, 0, 0, 0))
        ImageDraw.Draw(im).ellipse((i * 6, 20, i * 6 + 40, 60), fill=(255, 190, 0, 255))
        frames.append(im)
    frames[0].save(tmp / "spin.gif", save_all=True, append_images=frames[1:], duration=80, loop=0, disposal=2)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(layout)
    pic = s1.shapes.add_picture(str(tmp / "spin.gif"), 0, 0, Emu(900000), Emu(900000))
    _, rid = layout.part.get_or_add_image_part(str(tmp / "spin.gif"))
    pic._element.getparent().remove(pic._element)
    pic._element.xpath(".//a:blip")[0].set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
    pic._element.xpath(".//p:cNvPr")[0].set("id", "90")
    layout.shapes._spTree.append(pic._element)
    s1.shapes.add_movie(str(tmp / "bg.mp4"), 0, 0, prs.slide_width, prs.slide_height, mime_type="video/mp4")
    box = s1.shapes.add_shape(1, Emu(1500000), Emu(2500000), Emu(4000000), Emu(1500000))
    box.text_frame.text = "%42"
    s2 = prs.slides.add_slide(layout)
    s2.shapes.add_shape(1, Emu(1500000), Emu(2500000), Emu(4000000), Emu(1500000)).text_frame.text = "Merhaba"
    prs.save(str(tmp / "media.pptx"))
    return tmp / "media.pptx"


@unittest.skipUnless(shutil.which("ffmpeg"), "ffmpeg yok")
class Media(unittest.TestCase):
    def test_collects_video_and_layout_gif(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = media_deck(tmp)
            media = collect(deck, Path(tmp))
            self.assertEqual(sorted((m["type"], m["layer"]) for m in media[1]), [("gif", "layout"), ("video", "slide")])
            self.assertEqual([(m["type"], m["layer"]) for m in media[2]], [("gif", "layout")])
            video = next(m for m in media[1] if m["type"] == "video")
            self.assertAlmostEqual(video["duration"], 3, delta=0.2)
            self.assertTrue(video["audio"])

    @unittest.skipUnless(POWERPOINT and shutil.which("node"), "PowerPoint/Node yok")
    def test_background_video_plays_in_scene(self):
        """Medya açıkken iki an arasındaki fark, medya kapalıyken olandan çok daha büyük olmalı."""
        from PIL import Image, ImageChops, ImageStat
        with tempfile.TemporaryDirectory() as tmp:
            deck, work = media_deck(tmp), Path(tmp) / "w"

            def change(media_on):
                run("analyze", deck, "-w", work)
                plan = json.loads((work / "plan.json").read_text(encoding="utf-8"))
                plan["slides"][0]["media"] = media_on
                (work / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
                run("preview", deck, "-w", work, "--slides", "1")
                stills = sorted((work / "scenes" / "001").glob("still-*.png"), key=lambda p: float(p.stem[6:]))
                a, b = (Image.open(p).convert("RGB") for p in (stills[0], stills[2]))
                for p in stills:
                    p.unlink()
                return sum(ImageStat.Stat(ImageChops.difference(a, b)).mean)

            on, off = change(True), change(False)
            self.assertGreater(on, 4 * off + 5, (on, off))
            scene = json.loads((work / "scenes" / "001" / "scene.json").read_text(encoding="utf-8"))
            self.assertEqual(scene["media"], [])
            # düzendeki GIF onsuz çizilmiş arka planın (bgm.png) üstünde oynar: ilk karesi altta kalmaz
            bg, bgm = (Image.open(work / "layers" / "001" / n).convert("RGB") for n in ("bg.png", "bgm.png"))
            box = ImageChops.difference(bg, bgm).getbbox()
            self.assertIsNotNone(box)
            self.assertLess(box[2], bg.width * 0.15)


if __name__ == "__main__":
    unittest.main()
